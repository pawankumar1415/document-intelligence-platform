"""API routes for AI Narrative Search."""
from __future__ import annotations

import logging
from typing import Annotated, Any

from fastapi import APIRouter, Header, HTTPException, UploadFile, File, Form
from fastapi.responses import JSONResponse

from backend.app.models.schemas import (
    AdminUserRecord,
    AdminUserUpdateRequest,
    AnalyticsDashboard,
    AuthLoginRequest,
    AuthRegisterRequest,
    AuthResponse,
    BatchScoreResponse,
    ChatRequest,
    ChatResponse,
    ColumnDetectionResponse,
    DomainProfile,
    DriftMetrics,
    EmbeddingCatalog,
    EmbeddingConfigRequest,
    ExtractedRowsResponse,
    ExtractedTextResponse,
    ExcelRowRecord,
    FinancialUploadResponse,
    IngestReferenceResponse,
    NarrativeScoreResult,
    ProviderCatalogResponse,
    RubricCreateRequest,
    RulesUploadResponse,
    StandardsStatus,
    RubricRecord,
    RubricSummary,
    ScoreRequest,
    SharePointExtractRequest,
    SharePointFile,
    SharePointFilesResponse,
    SharePointLibrary,
)
from backend.app.services import auth_service, persistence
from backend.app.services.batch_service import run_batch_score
from backend.app.services.narrative_scorer import run_score
from backend.app.services.reference_service import (
    delete_reference_file,
    ingest_reference_file,
    list_reference_files,
)
from backend.app.services.provider_catalog import get_provider_catalog
from backend.app.services.embedding_service import (
    embedding_configuration,
    set_runtime_embedding_config,
)
from backend.app.services.vector_store import vector_store_status

logger = logging.getLogger(__name__)

router = APIRouter(prefix="/api/v1")


# ── Health ────────────────────────────────────────────────────────────────────

@router.get("/health")
def health() -> dict:
    vs = vector_store_status()
    return {
        "status": "ok",
        "product": "AI Narrative Search",
        "vector_store": vs["configured"],
        "embedding_backend": vs["embedding_backend"],
    }


# ── Auth ──────────────────────────────────────────────────────────────────────

@router.post("/auth/register", response_model=AuthResponse)
def register(body: AuthRegisterRequest) -> Any:
    email = body.email.lower().strip()
    if not email or not body.password:
        raise HTTPException(status_code=422, detail="Email and password are required.")
    if len(body.password) < 8:
        raise HTTPException(status_code=422, detail="Password must be at least 8 characters.")

    existing = persistence.get_user_by_email(email)
    if existing:
        raise HTTPException(status_code=409, detail="An account with this email already exists.")

    user = auth_service.create_user(email=email, password=body.password)
    token, expires_in = auth_service.create_user_session(user["id"])
    return {
        "access_token": token,
        "token_type": "bearer",
        "expires_in": expires_in,
        "user": {
            "id": user["id"],
            "email": user["email"],
            "is_admin": user["is_admin"],
            "is_active": user.get("is_active", True),
            "onboarding_completed": bool(user.get("onboarding_completed", False)),
        },
    }


@router.post("/auth/login", response_model=AuthResponse)
def login(body: AuthLoginRequest) -> Any:
    user = auth_service.verify_credentials(body.email, body.password)
    if not user:
        raise HTTPException(status_code=401, detail="Invalid email or password.")
    token, expires_in = auth_service.create_user_session(user["id"])
    return {
        "access_token": token,
        "token_type": "bearer",
        "expires_in": expires_in,
        "user": {
            "id": user["id"],
            "email": user["email"],
            "is_admin": user["is_admin"],
            "is_active": user.get("is_active", True),
            "onboarding_completed": bool(user.get("onboarding_completed", False)),
        },
    }


# ── Reference Files ───────────────────────────────────────────────────────────

@router.get("/references")
def list_references(authorization: Annotated[str | None, Header()] = None) -> Any:
    user = auth_service.require_authenticated_user(authorization)
    return list_reference_files(user["id"])


@router.get("/references/periods")
def get_reference_periods(authorization: Annotated[str | None, Header()] = None) -> Any:
    user = auth_service.require_authenticated_user(authorization)
    from backend.app.services.vector_store import get_available_periods
    try:
        return {"periods": get_available_periods(user["id"])}
    except Exception:
        return {"periods": []}


@router.get("/debug/vector-store")
def debug_vector_store(authorization: Annotated[str | None, Header()] = None) -> Any:
    user = auth_service.require_authenticated_user(authorization)
    from backend.app.services.vector_store import _connect
    from backend.app.services.embedding_service import (
        configured_embedding_backend, configured_embedding_model_id, embedding_configuration
    )
    result: dict[str, Any] = {
        "user_id": user["id"],
        "embedding_backend": configured_embedding_backend(),
        "embedding_model_id": configured_embedding_model_id(),
        "embedding_config": embedding_configuration(),
    }
    try:
        with _connect() as conn:
            with conn.cursor() as cur:
                cur.execute("SELECT COUNT(*) FROM reference_narratives WHERE user_id = %s;", (user["id"],))
                result["records_for_user"] = cur.fetchone()[0]
                cur.execute("SELECT COUNT(*) FROM reference_narratives;")
                result["total_records"] = cur.fetchone()[0]
                cur.execute("SELECT DISTINCT user_id FROM reference_narratives;")
                result["all_user_ids"] = [r[0] for r in cur.fetchall()]
                cur.execute(
                    "SELECT unique_id FROM reference_narratives WHERE user_id = %s LIMIT 5;",
                    (user["id"],)
                )
                result["sample_unique_ids"] = [r[0] for r in cur.fetchall()]
    except Exception as exc:
        result["db_error"] = str(exc)
    return result


@router.post("/references/ingest", response_model=IngestReferenceResponse)
async def ingest_reference(
    file: UploadFile = File(...),
    description: str = Form(default=""),
    authorization: Annotated[str | None, Header()] = None,
) -> Any:
    user = auth_service.require_authenticated_user(authorization)
    if not file.filename:
        raise HTTPException(status_code=422, detail="No file provided.")

    allowed_extensions = (".xlsx", ".xls", ".csv")
    if not any(file.filename.lower().endswith(ext) for ext in allowed_extensions):
        raise HTTPException(
            status_code=422,
            detail=f"Unsupported file type. Accepted: {', '.join(allowed_extensions)}",
        )

    raw_bytes = await file.read()
    if not raw_bytes:
        raise HTTPException(status_code=422, detail="Uploaded file is empty.")

    try:
        result = ingest_reference_file(
            raw_bytes=raw_bytes,
            filename=file.filename,
            user_id=user["id"],
            description=description,
        )
    except ValueError as exc:
        raise HTTPException(status_code=422, detail=str(exc))
    except Exception as exc:
        logger.error("Reference ingest failed: %s", exc, exc_info=True)
        raise HTTPException(status_code=500, detail=f"Ingest failed: {exc}")

    return result


@router.delete("/references/{file_id}")
def remove_reference(
    file_id: int,
    authorization: Annotated[str | None, Header()] = None,
) -> Any:
    user = auth_service.require_authenticated_user(authorization)
    deleted = delete_reference_file(file_id, user["id"])
    if not deleted:
        raise HTTPException(status_code=404, detail="Reference file not found.")
    return {"status": "deleted", "file_id": file_id}


# ── Domain Profile ────────────────────────────────────────────────────────────

@router.get("/domain/profile", response_model=DomainProfile)
def get_domain_profile(authorization: Annotated[str | None, Header()] = None) -> Any:
    user = auth_service.require_authenticated_user(authorization)
    profile = persistence.get_domain_profile(user["id"])
    return profile or {}


@router.delete("/domain/profile")
def delete_domain_profile(authorization: Annotated[str | None, Header()] = None) -> Any:
    user = auth_service.require_authenticated_user(authorization)
    persistence.delete_domain_profile(user["id"])
    return {"status": "deleted"}


# ── Single Narrative Scoring ──────────────────────────────────────────────────

@router.post("/score", response_model=NarrativeScoreResult)
def score_narrative(
    body: ScoreRequest,
    authorization: Annotated[str | None, Header()] = None,
) -> Any:
    user = auth_service.require_authenticated_user(authorization)
    try:
        result = run_score(
            narrative=body.narrative,
            unique_id=body.unique_id,
            document_name=body.document_name,
            user_id=user["id"],
            rubric_id=body.rubric_id,
            provider=body.llm_provider,
            model=body.llm_model,
            top_k_references=body.top_k_references,
        )
    except RuntimeError as exc:
        raise HTTPException(status_code=500, detail=str(exc))
    except Exception as exc:
        logger.error("Scoring failed: %s", exc, exc_info=True)
        raise HTTPException(status_code=500, detail=f"Scoring failed: {exc}")
    return result


# ── Batch Scoring ─────────────────────────────────────────────────────────────

@router.post("/score/batch", response_model=BatchScoreResponse)
async def score_batch(
    file: UploadFile = File(...),
    rubric_id: int | None = Form(default=None),
    llm_provider: str = Form(default="openai"),
    llm_model: str | None = Form(default=None),
    top_k_references: int = Form(default=5),
    id_column: str | None = Form(default=None),
    narrative_column: str | None = Form(default=None),
    authorization: Annotated[str | None, Header()] = None,
) -> Any:
    user = auth_service.require_authenticated_user(authorization)
    if not file.filename:
        raise HTTPException(status_code=422, detail="No file provided.")

    raw_bytes = await file.read()
    if not raw_bytes:
        raise HTTPException(status_code=422, detail="Uploaded file is empty.")

    try:
        result = run_batch_score(
            raw_bytes=raw_bytes,
            filename=file.filename,
            user_id=user["id"],
            rubric_id=rubric_id,
            provider=llm_provider,  # type: ignore[arg-type]
            model=llm_model or None,
            top_k_references=top_k_references,
            id_column=id_column or None,
            narrative_column=narrative_column or None,
        )
    except ValueError as exc:
        raise HTTPException(status_code=422, detail=str(exc))
    except Exception as exc:
        logger.error("Batch scoring failed: %s", exc, exc_info=True)
        raise HTTPException(status_code=500, detail=f"Batch scoring failed: {exc}")

    return result


# ── Score History ─────────────────────────────────────────────────────────────

@router.get("/scores")
def list_scores(
    limit: int = 50,
    authorization: Annotated[str | None, Header()] = None,
) -> Any:
    user = auth_service.require_authenticated_user(authorization)
    return persistence.list_score_results(user["id"], limit=min(limit, 200))


@router.get("/scores/{result_id}")
def get_score(
    result_id: int,
    authorization: Annotated[str | None, Header()] = None,
) -> Any:
    user = auth_service.require_authenticated_user(authorization)
    record = persistence.get_score_result(result_id, user["id"])
    if not record:
        raise HTTPException(status_code=404, detail="Score result not found.")
    return record


# ── Rubrics ───────────────────────────────────────────────────────────────────

@router.get("/rubrics", response_model=list[RubricSummary])
def list_rubrics(authorization: Annotated[str | None, Header()] = None) -> Any:
    user = auth_service.require_authenticated_user(authorization)
    rubric_id = persistence.ensure_default_rubric(user["id"])
    _ = rubric_id
    return persistence.list_rubrics(user["id"])


@router.post("/rubrics", response_model=RubricRecord)
def create_rubric(
    body: RubricCreateRequest,
    authorization: Annotated[str | None, Header()] = None,
) -> Any:
    user = auth_service.require_authenticated_user(authorization)
    rubric_id = persistence.create_rubric(
        user_id=user["id"],
        name=body.name,
        description=body.description,
        criteria=[c.model_dump() for c in body.criteria],
    )
    rubric = persistence.get_rubric(rubric_id, user["id"])
    return {
        "id": rubric["id"],
        "name": rubric["name"],
        "description": rubric["description"],
        "criteria": rubric["criteria"],
        "is_default": bool(rubric["is_default"]),
    }


@router.delete("/rubrics/{rubric_id}")
def delete_rubric(
    rubric_id: int,
    authorization: Annotated[str | None, Header()] = None,
) -> Any:
    user = auth_service.require_authenticated_user(authorization)
    persistence.delete_rubric(rubric_id, user["id"])
    return {"status": "deleted"}


# ── Analytics ─────────────────────────────────────────────────────────────────

@router.get("/analytics", response_model=AnalyticsDashboard)
def analytics(authorization: Annotated[str | None, Header()] = None) -> Any:
    user = auth_service.require_authenticated_user(authorization)
    overview = persistence.get_analytics_overview(user["id"])
    recent = persistence.list_score_results(user["id"], limit=20)
    return {
        "overview": overview,
        "recent_scores": [
            {
                "id": r["id"],
                "unique_id": r["unique_id"],
                "document_name": r["document_name"],
                "overall_verdict": r["overall_verdict"],
                "compliance_score": r["compliance_score"],
                "created_at": r["created_at"],
            }
            for r in recent
        ],
    }


# ── Provider Catalog ──────────────────────────────────────────────────────────

@router.get("/providers/models", response_model=ProviderCatalogResponse)
def provider_models(authorization: Annotated[str | None, Header()] = None) -> Any:
    auth_service.require_authenticated_user(authorization)
    catalog = get_provider_catalog()
    return {"providers": catalog}


# ── Embedding Config ──────────────────────────────────────────────────────────

@router.get("/embedding/config")
def embedding_config(authorization: Annotated[str | None, Header()] = None) -> Any:
    auth_service.require_authenticated_user(authorization)
    return embedding_configuration()


@router.post("/embedding/config", response_model=EmbeddingCatalog)
def update_embedding(
    body: EmbeddingConfigRequest,
    authorization: Annotated[str | None, Header()] = None,
) -> Any:
    auth_service.require_authenticated_user(authorization)
    try:
        set_runtime_embedding_config(backend=body.backend, model_id=body.model_id)
    except RuntimeError as exc:
        raise HTTPException(status_code=422, detail=str(exc))

    # Re-initialise the vector store so a dimension mismatch is caught immediately.
    # If VECTOR_STORE_RESET_ON_MISMATCH=true the tables are rebuilt automatically;
    # otherwise an error is raised here rather than failing silently at ingest time.
    from backend.app.services.vector_store import init_vector_store, vector_store_status
    if vector_store_status()["configured"]:
        try:
            init_vector_store()
        except RuntimeError as exc:
            raise HTTPException(status_code=409, detail=str(exc))

    return embedding_configuration()


# ── Admin ─────────────────────────────────────────────────────────────────────

@router.get("/admin/users", response_model=list[AdminUserRecord])
def admin_list_users(authorization: Annotated[str | None, Header()] = None) -> Any:
    auth_service.require_admin_user(authorization)
    users = persistence.list_all_users()
    return [
        {
            "id": u["id"],
            "email": u["email"],
            "is_admin": bool(u["is_admin"]),
            "is_active": bool(u["is_active"]),
            "created_at": u["created_at"],
            "score_count": int(u.get("score_count", 0)),
        }
        for u in users
    ]


@router.patch("/admin/users/{user_id}", response_model=AdminUserRecord)
def admin_update_user(
    user_id: int,
    body: AdminUserUpdateRequest,
    authorization: Annotated[str | None, Header()] = None,
) -> Any:
    auth_service.require_admin_user(authorization)
    persistence.update_admin_flags(user_id, is_admin=body.is_admin, is_active=body.is_active)
    users = persistence.list_all_users()
    user = next((u for u in users if u["id"] == user_id), None)
    if not user:
        raise HTTPException(status_code=404, detail="User not found.")
    return {
        "id": user["id"],
        "email": user["email"],
        "is_admin": bool(user["is_admin"]),
        "is_active": bool(user["is_active"]),
        "created_at": user["created_at"],
        "score_count": int(user.get("score_count", 0)),
    }


@router.delete("/admin/users/{user_id}")
def admin_delete_user(
    user_id: int,
    authorization: Annotated[str | None, Header()] = None,
) -> Any:
    admin = auth_service.require_admin_user(authorization)
    if admin["id"] == user_id:
        raise HTTPException(status_code=400, detail="Cannot delete your own account.")
    persistence.delete_user(user_id)
    return {"status": "deleted", "user_id": user_id}


# ── Onboarding ────────────────────────────────────────────────────────────────

@router.post("/user/complete-onboarding")
def complete_onboarding(
    authorization: Annotated[str | None, Header()] = None,
) -> Any:
    user = auth_service.require_authenticated_user(authorization)
    persistence.complete_onboarding(int(user["id"]))
    return {"status": "ok"}


# ── File text extraction ──────────────────────────────────────────────────────

@router.post("/file/extract-text", response_model=ExtractedTextResponse)
async def extract_text_from_upload(
    file: UploadFile = File(...),
    authorization: Annotated[str | None, Header()] = None,
) -> ExtractedTextResponse:
    auth_service.require_authenticated_user(authorization)
    from backend.app.services.text_extractor import extract_text
    content = await file.read()
    try:
        text = extract_text(file.filename or "", content)
    except (ValueError, RuntimeError) as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc
    return ExtractedTextResponse(
        filename=file.filename or "",
        text=text,
        char_count=len(text),
    )


@router.post("/file/extract-rows", response_model=ExtractedRowsResponse)
async def extract_rows_from_upload(
    file: UploadFile = File(...),
    authorization: Annotated[str | None, Header()] = None,
) -> ExtractedRowsResponse:
    auth_service.require_authenticated_user(authorization)
    from backend.app.services.excel_parser import (
        parse_excel_for_scoring, parse_csv_for_scoring,
        parse_docx_for_scoring, parse_pdf_for_scoring,
        _ID_KEYWORDS, _TEXT_KEYWORDS, _find_column,
    )
    import io, pandas as pd

    content = await file.read()
    fname = (file.filename or "").lower()
    ext = fname.rsplit(".", 1)[-1] if "." in fname else ""

    try:
        if ext == "csv":
            records = parse_csv_for_scoring(content)
            df = pd.read_csv(io.BytesIO(content), dtype=str)
            col_map = {str(c).strip().lower(): str(c) for c in df.columns}
            id_col = _find_column(col_map, _ID_KEYWORDS) or str(df.columns[0])
            txt_col = _find_column(col_map, _TEXT_KEYWORDS) or str(df.columns[-1])
        elif ext == "docx":
            records = parse_docx_for_scoring(content)
            from docx import Document as _DocxDoc  # type: ignore
            _doc = _DocxDoc(io.BytesIO(content))
            headers = [c.text.strip() for c in _doc.tables[0].rows[0].cells] if _doc.tables else []
            col_map = {h.lower(): h for h in headers}
            id_col = _find_column(col_map, _ID_KEYWORDS) or (headers[0] if headers else "id")
            txt_col = _find_column(col_map, _TEXT_KEYWORDS) or (headers[-1] if headers else "text")
        elif ext == "pdf":
            records = parse_pdf_for_scoring(content)
            import pdfplumber  # type: ignore
            with pdfplumber.open(io.BytesIO(content)) as _pdf:
                headers = []
                for _page in _pdf.pages:
                    for _tbl in _page.extract_tables():
                        if _tbl:
                            headers = [str(c or "").strip() for c in _tbl[0]]
                            break
                    if headers:
                        break
            col_map = {h.lower(): h for h in headers}
            id_col = _find_column(col_map, _ID_KEYWORDS) or (headers[0] if headers else "id")
            txt_col = _find_column(col_map, _TEXT_KEYWORDS) or (headers[-1] if headers else "text")
        else:
            records = parse_excel_for_scoring(content, filename=file.filename or "")
            df = pd.read_excel(io.BytesIO(content), header=0, dtype=str)
            col_map = {str(c).strip().lower(): str(c) for c in df.columns}
            id_col = _find_column(col_map, _ID_KEYWORDS) or str(df.columns[0])
            txt_col = _find_column(col_map, _TEXT_KEYWORDS) or str(df.columns[-1])
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc

    rows = [
        ExcelRowRecord(id=r["unique_id"], narrative=r["narrative_text"])
        for r in records
        if r.get("unique_id") and r.get("narrative_text")
    ]
    return ExtractedRowsResponse(
        filename=file.filename or "",
        id_column=id_col,
        narrative_column=txt_col,
        rows=rows,
    )


# ── SharePoint ────────────────────────────────────────────────────────────────

@router.get("/sharepoint/status")
def sharepoint_status(
    authorization: Annotated[str | None, Header()] = None,
) -> Any:
    auth_service.require_authenticated_user(authorization)
    from backend.app.services.sharepoint_service import is_configured
    return {"configured": is_configured()}


@router.get("/sharepoint/libraries", response_model=list[SharePointLibrary])
def get_sharepoint_libraries(
    site_id: str | None = None,
    authorization: Annotated[str | None, Header()] = None,
) -> list[SharePointLibrary]:
    auth_service.require_authenticated_user(authorization)
    from backend.app.services.sharepoint_service import list_libraries
    try:
        libs = list_libraries(site_id=site_id)
    except RuntimeError as exc:
        raise HTTPException(status_code=503, detail=str(exc)) from exc
    return [SharePointLibrary(**lib) for lib in libs]


@router.get("/sharepoint/files", response_model=SharePointFilesResponse)
def get_sharepoint_files(
    library_id: str,
    folder_path: str = "/",
    authorization: Annotated[str | None, Header()] = None,
) -> SharePointFilesResponse:
    auth_service.require_authenticated_user(authorization)
    from backend.app.services.sharepoint_service import list_files
    try:
        items = list_files(library_id=library_id, folder_path=folder_path)
    except RuntimeError as exc:
        raise HTTPException(status_code=503, detail=str(exc)) from exc
    return SharePointFilesResponse(
        library_id=library_id,
        folder_path=folder_path,
        items=[SharePointFile(**item) for item in items],
    )


@router.get("/sharepoint/download")
def download_sharepoint_file(
    library_id: str,
    item_id: str,
    filename: str,
    authorization: Annotated[str | None, Header()] = None,
) -> Any:
    auth_service.require_authenticated_user(authorization)
    from backend.app.services.sharepoint_service import download_file
    from fastapi.responses import Response as FastResponse

    try:
        content = download_file(library_id=library_id, item_id=item_id)
    except RuntimeError as exc:
        raise HTTPException(status_code=503, detail=str(exc)) from exc

    ext = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    media_type = (
        "text/csv"
        if ext == "csv"
        else "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        if ext in ("xlsx", "xls")
        else "application/octet-stream"
    )
    return FastResponse(
        content=content,
        media_type=media_type,
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


@router.post("/sharepoint/extract-text", response_model=ExtractedTextResponse)
def sharepoint_extract_text(
    request: SharePointExtractRequest,
    authorization: Annotated[str | None, Header()] = None,
) -> ExtractedTextResponse:
    auth_service.require_authenticated_user(authorization)
    from backend.app.services.sharepoint_service import download_file
    from backend.app.services.text_extractor import extract_text
    try:
        content = download_file(library_id=request.library_id, item_id=request.item_id)
    except RuntimeError as exc:
        raise HTTPException(status_code=503, detail=str(exc)) from exc
    try:
        text = extract_text(request.filename, content)
    except (ValueError, RuntimeError) as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc
    return ExtractedTextResponse(
        filename=request.filename,
        text=text,
        char_count=len(text),
    )


@router.post("/sharepoint/extract-rows", response_model=ExtractedRowsResponse)
def sharepoint_extract_rows(
    request: SharePointExtractRequest,
    authorization: Annotated[str | None, Header()] = None,
) -> ExtractedRowsResponse:
    auth_service.require_authenticated_user(authorization)
    from backend.app.services.sharepoint_service import download_file
    from backend.app.services.excel_parser import (
        parse_excel_for_scoring,
        parse_csv_for_scoring,
        parse_docx_for_scoring,
        parse_pdf_for_scoring,
        _ID_KEYWORDS,
        _TEXT_KEYWORDS,
        _find_column,
    )
    import io
    import pandas as pd

    try:
        content = download_file(library_id=request.library_id, item_id=request.item_id)
    except RuntimeError as exc:
        raise HTTPException(status_code=503, detail=str(exc)) from exc

    fname = request.filename.lower()
    try:
        if fname.endswith(".csv"):
            records = parse_csv_for_scoring(content)
            df = pd.read_csv(io.BytesIO(content), dtype=str)
            col_map = {str(c).strip().lower(): str(c) for c in df.columns}
            id_col = _find_column(col_map, _ID_KEYWORDS) or str(df.columns[0])
            txt_col = _find_column(col_map, _TEXT_KEYWORDS) or str(df.columns[-1])
        elif fname.endswith(".docx"):
            records = parse_docx_for_scoring(content)
            from docx import Document as _DocxDoc  # type: ignore
            _doc = _DocxDoc(io.BytesIO(content))
            headers = [c.text.strip() for c in _doc.tables[0].rows[0].cells] if _doc.tables else []
            col_map = {h.lower(): h for h in headers}
            id_col = _find_column(col_map, _ID_KEYWORDS) or (headers[0] if headers else "id")
            txt_col = _find_column(col_map, _TEXT_KEYWORDS) or (headers[-1] if headers else "text")
        elif fname.endswith(".pdf"):
            records = parse_pdf_for_scoring(content)
            import pdfplumber  # type: ignore
            with pdfplumber.open(io.BytesIO(content)) as _pdf:
                headers = []
                for _page in _pdf.pages:
                    for _tbl in _page.extract_tables():
                        if _tbl:
                            headers = [str(c or "").strip() for c in _tbl[0]]
                            break
                    if headers:
                        break
            col_map = {h.lower(): h for h in headers}
            id_col = _find_column(col_map, _ID_KEYWORDS) or (headers[0] if headers else "id")
            txt_col = _find_column(col_map, _TEXT_KEYWORDS) or (headers[-1] if headers else "text")
        else:
            records = parse_excel_for_scoring(content, filename=request.filename)
            df = pd.read_excel(io.BytesIO(content), header=0, dtype=str)
            col_map = {str(c).strip().lower(): str(c) for c in df.columns}
            id_col = _find_column(col_map, _ID_KEYWORDS) or str(df.columns[0])
            txt_col = _find_column(col_map, _TEXT_KEYWORDS) or str(df.columns[-1])
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc

    rows = [
        ExcelRowRecord(id=r["unique_id"], narrative=r["narrative_text"])
        for r in records
        if r.get("unique_id") and r.get("narrative_text")
    ]
    return ExtractedRowsResponse(
        filename=request.filename,
        id_column=id_col,
        narrative_column=txt_col,
        rows=rows,
    )


# ── Chat (RAG over user's reference library) ─────────────────────────────────

@router.post("/chat", response_model=ChatResponse)
def chat(
    body: ChatRequest,
    authorization: Annotated[str | None, Header()] = None,
) -> Any:
    user = auth_service.require_authenticated_user(authorization)
    from backend.app.services.chat_service import run_chat
    from backend.app.services.llm_provider import LLMProvider

    from backend.app.config import default_llm_provider
    provider: LLMProvider = body.provider if body.provider in ("openai", "groq", "azure_openai", "ollama") else default_llm_provider()  # type: ignore[assignment]
    try:
        result = run_chat(
            user_id=user["id"],
            messages=[{"role": m.role, "content": m.content} for m in body.messages],
            provider=provider,
            model=body.model,
            top_k=body.top_k,
        )
    except RuntimeError as exc:
        raise HTTPException(status_code=502, detail=str(exc)) from exc
    return ChatResponse(
        reply=result["reply"],
        sources=[
            {"unique_id": s["unique_id"], "excerpt": s["excerpt"], "score": s["score"]}
            for s in result["sources"]
        ],
    )


# ── Column Detection ──────────────────────────────────────────────────────────

@router.post("/columns/detect", response_model=ColumnDetectionResponse)
async def detect_columns(
    file: UploadFile = File(...),
    use_llm: bool = False,
    llm_provider: str | None = None,
    llm_model: str | None = None,
    authorization: Annotated[str | None, Header()] = None,
) -> ColumnDetectionResponse:
    auth_service.require_authenticated_user(authorization)
    import io
    import pandas as pd
    from backend.app.models.schemas import ColumnCandidate

    content = await file.read()
    fname = (file.filename or "").lower()
    ext = fname.rsplit(".", 1)[-1] if "." in fname else ""

    try:
        if ext == "csv":
            df = pd.read_csv(io.BytesIO(content), nrows=20, dtype=str)
        elif ext == "docx":
            from docx import Document as _DocxDoc  # type: ignore
            _doc = _DocxDoc(io.BytesIO(content))
            if not _doc.tables:
                raise ValueError("No tables found in .docx file.")
            _tbl = _doc.tables[0]
            _headers = [c.text.strip() for c in _tbl.rows[0].cells]
            _data = [[c.text.strip() for c in r.cells] for r in _tbl.rows[1:21]]
            df = pd.DataFrame(_data, columns=_headers, dtype=str)
        elif ext == "pdf":
            import pdfplumber  # type: ignore
            _rows_all: list[list[str]] = []
            with pdfplumber.open(io.BytesIO(content)) as _pdf:
                for _page in _pdf.pages:
                    for _t in _page.extract_tables():
                        _rows_all.extend([[str(c or "").strip() for c in r] for r in _t])
                        if len(_rows_all) > 21:
                            break
                    if len(_rows_all) > 21:
                        break
            if len(_rows_all) < 2:
                raise ValueError("No tables found in PDF file.")
            df = pd.DataFrame(_rows_all[1:21], columns=_rows_all[0], dtype=str)
        else:
            df = pd.read_excel(io.BytesIO(content), nrows=20, dtype=str)
    except Exception as exc:
        raise HTTPException(status_code=400, detail=f"Cannot parse file: {exc}") from exc

    all_columns = list(df.columns)
    if not all_columns:
        raise HTTPException(status_code=400, detail="File has no columns.")

    _ID_KWS = {"id", "uid", "name", "code", "ref", "reference", "identifier", "document", "key", "no"}
    _TEXT_KWS = {"narrative", "text", "description", "content", "body", "summary", "commentary", "note", "comment"}

    def _last_word(col: str) -> str:
        import re as _re
        parts = _re.split(r"[^a-z0-9]", col.lower())
        return parts[-1] if parts else ""

    def score_id(col: str) -> float:
        vals = df[col].dropna().astype(str)
        if len(vals) == 0:
            return 0.0
        avg_len = vals.str.len().mean()
        uniqueness = vals.nunique() / len(vals)
        # Allow longer IDs (project names can be 50+ chars); use 100-char cutoff
        length_score = max(0.0, 1.0 - avg_len / 100.0)
        # Boost when the last word of the column name signals an identifier
        keyword_bonus = 0.35 if _last_word(col) in _ID_KWS else 0.0
        return round(uniqueness * 0.55 + length_score * 0.1 + keyword_bonus, 3)

    def score_narrative(col: str) -> float:
        vals = df[col].dropna().astype(str)
        if len(vals) == 0:
            return 0.0
        avg_len = vals.str.len().mean()
        # Good narrative: long text
        length_score = min(avg_len / 300.0, 1.0)
        keyword_bonus = 0.2 if _last_word(col) in _TEXT_KWS else 0.0
        return round(length_score * 0.8 + keyword_bonus, 3)

    def sample_vals(col: str) -> list[str]:
        return [str(v)[:80] for v in df[col].dropna().head(3).tolist()]

    id_scored = sorted(
        [(c, score_id(c)) for c in all_columns], key=lambda x: -x[1]
    )
    narrative_scored = sorted(
        [(c, score_narrative(c)) for c in all_columns], key=lambda x: -x[1]
    )

    best_id = id_scored[0][0] if id_scored else None
    best_narrative = next(
        (c for c, _ in narrative_scored if c != best_id), None
    )

    id_candidates = [
        ColumnCandidate(name=c, confidence=s, sample=sample_vals(c))
        for c, s in id_scored[:5]
    ]
    narrative_candidates = [
        ColumnCandidate(name=c, confidence=s, sample=sample_vals(c))
        for c, s in narrative_scored[:5] if c != best_id
    ]

    ambiguous = (
        len(id_scored) >= 2 and abs(id_scored[0][1] - id_scored[1][1]) < 0.1
    ) or (
        len(narrative_scored) >= 2
        and abs(narrative_scored[0][1] - narrative_scored[1][1]) < 0.1
    )

    method = "heuristic"

    if use_llm and (ambiguous or True):
        try:
            from backend.app.config import default_llm_provider
            from backend.app.services.llm_provider import generate_json_object

            provider = llm_provider if llm_provider in ("openai", "groq", "azure_openai", "ollama") else default_llm_provider()
            sample_rows = df.head(3).fillna("").to_dict(orient="records")
            prompt = (
                "You are analysing a spreadsheet to identify which column contains a unique record ID "
                "and which column contains the main narrative or description text.\n\n"
                f"Column names: {all_columns}\n\n"
                f"First 3 rows of data:\n{sample_rows}\n\n"
                "Return JSON with exactly these keys:\n"
                '{"id_column": "<column name or null>", "narrative_column": "<column name or null>"}\n'
                "Only return column names that exist in the list above. Return null if unsure."
            )
            result = generate_json_object(
                system_prompt="You identify spreadsheet columns by their content.",
                user_prompt=prompt,
                provider=provider,
                model=llm_model,
            )
            llm_id = result.get("id_column")
            llm_narrative = result.get("narrative_column")
            if llm_id in all_columns:
                best_id = llm_id
            if llm_narrative in all_columns and llm_narrative != best_id:
                best_narrative = llm_narrative
            method = "llm"
        except Exception as exc:
            logger.warning("LLM column detection failed, using heuristic result: %s", exc)

    return ColumnDetectionResponse(
        all_columns=all_columns,
        id_column=best_id,
        narrative_column=best_narrative,
        id_candidates=id_candidates,
        narrative_candidates=narrative_candidates,
        method=method,
        ambiguous=ambiguous,
    )


# ── Standards: Rules Document ─────────────────────────────────────────────────

@router.post("/standards/rules/upload", response_model=RulesUploadResponse)
async def upload_rules_document(
    file: UploadFile = File(...),
    authorization: Annotated[str | None, Header()] = None,
) -> Any:
    user = auth_service.require_authenticated_user(authorization)
    raw = await file.read()
    filename = file.filename or "rules"
    from backend.app.config import default_llm_provider
    from backend.app.services.rules_parser import parse_rules_document
    try:
        criteria = parse_rules_document(raw, filename, provider=default_llm_provider())
    except Exception as exc:
        raise HTTPException(status_code=422, detail=f"Could not parse rules document: {exc}")
    if not criteria:
        raise HTTPException(status_code=422, detail="No compliance criteria found in the uploaded document.")

    rubric_id = persistence.create_rubric(
        user_id=user["id"],
        name=f"Uploaded Rules — {filename}",
        description=f"Auto-extracted from {filename}",
        criteria=criteria,
    )
    persistence.set_active_rules_rubric(user["id"], rubric_id)
    return RulesUploadResponse(
        status="ok",
        rubric_id=rubric_id,
        criteria_count=len(criteria),
        filename=filename,
        message=f"{len(criteria)} criteria extracted and set as active rules.",
    )


@router.get("/standards/rules", response_model=StandardsStatus)
def get_standards_status(
    authorization: Annotated[str | None, Header()] = None,
) -> Any:
    user = auth_service.require_authenticated_user(authorization)
    uid = user["id"]

    active_rubric_id = persistence.get_active_rules_rubric_id(uid)
    if active_rubric_id:
        rubric = persistence.get_rubric(active_rubric_id, uid)
        rules_status = {
            "active": True,
            "rubric_id": active_rubric_id,
            "rubric_name": rubric["name"] if rubric else None,
            "criteria_count": len(rubric["criteria"]) if rubric else 0,
            "source_filename": rubric["name"].replace("Uploaded Rules — ", "") if rubric else None,
        }
    else:
        rules_status = {"active": False, "rubric_id": None, "rubric_name": None, "criteria_count": 0, "source_filename": None}

    fin = persistence.get_financial_upload_status(uid)
    financial_status = {
        "active": fin is not None,
        "filename": fin["filename"] if fin else None,
        "record_count": fin["record_count"] if fin else 0,
        "uploaded_at": fin["uploaded_at"] if fin else None,
    }

    return StandardsStatus(
        rules=rules_status,
        financial=financial_status,
    )


@router.delete("/standards/rules")
def delete_rules_document(authorization: Annotated[str | None, Header()] = None) -> Any:
    user = auth_service.require_authenticated_user(authorization)
    persistence.clear_active_rules_rubric(user["id"])
    return {"status": "deleted"}


# ── Standards: Financial Data ─────────────────────────────────────────────────

@router.post("/standards/financial/upload", response_model=FinancialUploadResponse)
async def upload_financial_data(
    file: UploadFile = File(...),
    authorization: Annotated[str | None, Header()] = None,
) -> Any:
    user = auth_service.require_authenticated_user(authorization)
    raw = await file.read()
    filename = file.filename or "financial_data"
    from backend.app.services.financial_service import ingest_financial_file
    try:
        result = ingest_financial_file(raw, filename, user_id=user["id"])
    except RuntimeError as exc:
        raise HTTPException(status_code=422, detail=str(exc))
    return FinancialUploadResponse(**result)


@router.delete("/standards/financial")
def delete_financial_data(authorization: Annotated[str | None, Header()] = None) -> Any:
    user = auth_service.require_authenticated_user(authorization)
    persistence.delete_financial_upload(user["id"])
    return {"status": "deleted"}


# ── Analytics: Drift Dashboard ────────────────────────────────────────────────

@router.get("/analytics/drift", response_model=DriftMetrics)
def get_drift_metrics(
    days: int = 30,
    authorization: Annotated[str | None, Header()] = None,
) -> Any:
    user = auth_service.require_authenticated_user(authorization)
    from backend.app.services.drift_service import get_drift_metrics
    return get_drift_metrics(user["id"], days=days)


@router.get("/analytics/drift/export")
def export_drift_csv(
    days: int = 30,
    authorization: Annotated[str | None, Header()] = None,
) -> Any:
    from fastapi.responses import Response
    user = auth_service.require_authenticated_user(authorization)
    from backend.app.services.drift_service import export_audit_csv
    csv_data = export_audit_csv(user["id"], days=days)
    return Response(
        content=csv_data,
        media_type="text/csv",
        headers={"Content-Disposition": f"attachment; filename=ai_audit_{days}d.csv"},
    )
