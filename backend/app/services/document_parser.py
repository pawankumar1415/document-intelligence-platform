from __future__ import annotations

from io import BytesIO
from pathlib import Path
from zipfile import BadZipFile
import re

from docx import Document
from fastapi import UploadFile

from backend.app.config import env, env_bool
from backend.app.models.schemas import ParsedDocument, ParsedSection
from backend.app.services.extraction_signals import extract_signals


_IMAGE_SUFFIXES = {".jpg", ".jpeg", ".png", ".webp", ".bmp", ".tiff", ".tif"}
_IMAGE_MIME: dict[str, str] = {
    ".jpg": "image/jpeg", ".jpeg": "image/jpeg",
    ".png": "image/png",  ".webp": "image/webp",
    ".bmp": "image/bmp",  ".tiff": "image/tiff", ".tif": "image/tiff",
}
_EXCEL_SUFFIXES = {".xlsx", ".xls"}

_IMAGE_EXTRACT_PROMPT = """\
You are an expert document analyst. Extract ALL text and information from this image.

If the image contains a document, letter, or report: extract every visible word, preserving headings and structure.
If it contains a table or spreadsheet: reproduce it as a markdown table.
If it contains a chart or graph: describe the chart type, title, axis labels, legend, and key data values.
If it contains a diagram or flowchart: describe the structure, steps, and labels in full.
If it contains handwriting: transcribe it as accurately as possible.

Return only the extracted content, well-structured. Do not add commentary or preamble."""


class DocumentParser:
    async def parse_upload(
        self,
        upload_file: UploadFile,
        llm_provider: str = "openai",
        llm_model: str | None = None,
    ) -> ParsedDocument:
        suffix = Path(upload_file.filename or "").suffix.lower()
        raw_bytes = await upload_file.read()

        if suffix == ".txt":
            return self._parse_txt(upload_file.filename or "document.txt", raw_bytes)
        if suffix == ".docx":
            return self._parse_docx(upload_file.filename or "document.docx", raw_bytes)
        if suffix == ".pdf":
            return self._parse_pdf_docling(upload_file.filename or "document.pdf", raw_bytes)
        if suffix in _IMAGE_SUFFIXES:
            return self._parse_image(
                upload_file.filename or f"image{suffix}", raw_bytes,
                llm_provider=llm_provider, llm_model=llm_model,
            )
        if suffix in _EXCEL_SUFFIXES:
            return self._parse_excel(upload_file.filename or "spreadsheet.xlsx", raw_bytes)
        if suffix == ".csv":
            return self._parse_csv(upload_file.filename or "data.csv", raw_bytes)

        raise ValueError(
            f"Unsupported file type '{suffix}'. "
            "Supported formats: .txt, .docx, .pdf, .xlsx, .xls, .csv, "
            ".jpg, .jpeg, .png, .webp, .bmp, .tiff"
        )

    def _parse_txt(self, filename: str, raw_bytes: bytes) -> ParsedDocument:
        text = raw_bytes.decode("utf-8", errors="ignore").replace("\r\n", "\n")
        paragraphs = [item.strip() for item in text.split("\n\n") if item.strip()]
        sections = [ParsedSection(heading=f"Section {index}", body=paragraph) for index, paragraph in enumerate(paragraphs, start=1)]
        return self._build_document(filename=filename, file_type="txt", title=Path(filename).stem, text=text, sections=sections)

    def _parse_docx(self, filename: str, raw_bytes: bytes) -> ParsedDocument:
        try:
            document = Document(BytesIO(raw_bytes))
        except BadZipFile as exc:
            raise ValueError("Uploaded .docx file is invalid or corrupted.") from exc

        sections: list[ParsedSection] = []
        full_paragraphs: list[str] = []
        current_heading = Path(filename).stem
        current_body: list[str] = []

        for paragraph in document.paragraphs:
            text = paragraph.text.strip()
            if not text:
                continue

            full_paragraphs.append(text)
            style_name = paragraph.style.name.lower() if paragraph.style and paragraph.style.name else ""

            if style_name.startswith("heading"):
                if current_body:
                    sections.append(ParsedSection(heading=current_heading, body="\n".join(current_body)))
                    current_body = []
                current_heading = text
                continue

            current_body.append(text)

        if current_body:
            sections.append(ParsedSection(heading=current_heading, body="\n".join(current_body)))

        if not sections and full_paragraphs:
            sections = [ParsedSection(heading=Path(filename).stem, body="\n".join(full_paragraphs))]

        return self._build_document(
            filename=filename,
            file_type="docx",
            title=Path(filename).stem,
            text="\n".join(full_paragraphs),
            sections=sections,
        )

    def _parse_pdf_docling(self, filename: str, raw_bytes: bytes) -> ParsedDocument:
        try:
            from docling.datamodel.base_models import DocumentStream, InputFormat
            from docling.datamodel.pipeline_options import PdfPipelineOptions, RapidOcrOptions, TableStructureOptions
            from docling.document_converter import DocumentConverter, PdfFormatOption
        except ModuleNotFoundError as exc:
            raise RuntimeError(
                "Docling is not installed. Install dependencies with `pip install -r requirements.txt`."
            ) from exc

        pipeline_options = PdfPipelineOptions()
        pipeline_options.do_ocr = True
        pipeline_options.do_table_structure = True
        pipeline_options.table_structure_options = TableStructureOptions(do_cell_matching=True)
        pipeline_options.force_full_page_ocr = env_bool("DOCLING_FORCE_FULL_PAGE_OCR", False)

        ocr_engine = (env("DOCLING_OCR_ENGINE", "rapidocr") or "rapidocr").lower().strip()
        if ocr_engine in {"rapidocr", "rapid_ocr"}:
            pipeline_options.ocr_options = RapidOcrOptions()
        elif ocr_engine == "tesseract_cli":
            try:
                from docling.datamodel.pipeline_options import TesseractCliOcrOptions

                pipeline_options.ocr_options = TesseractCliOcrOptions()
            except ModuleNotFoundError:
                pipeline_options.ocr_options = RapidOcrOptions()
        else:
            pipeline_options.ocr_options = RapidOcrOptions()

        converter = DocumentConverter(
            format_options={
                InputFormat.PDF: PdfFormatOption(pipeline_options=pipeline_options),
            }
        )
        source = DocumentStream(name=filename, stream=BytesIO(raw_bytes))

        try:
            result = converter.convert(source)
        except Exception as exc:
            raise ValueError(f"Docling could not parse this PDF: {exc}") from exc

        text = self._extract_docling_text(result)
        if not text.strip():
            raise ValueError("Docling parsed the PDF but no extractable text was found.")

        sections = self._sections_from_markdown_or_text(text, fallback_heading=Path(filename).stem)
        return self._build_document(
            filename=filename,
            file_type="pdf",
            title=Path(filename).stem,
            text=text,
            sections=sections,
        )

    def _build_document(
        self,
        *,
        filename: str,
        file_type: str,
        title: str,
        text: str,
        sections: list[ParsedSection],
    ) -> ParsedDocument:
        paragraph_count = sum(1 for block in text.splitlines() if block.strip())
        word_count = len(text.split())
        return ParsedDocument(
            filename=filename,
            file_type=file_type,
            title=title,
            text=text.strip(),
            sections=sections,
            word_count=word_count,
            paragraph_count=paragraph_count,
            extraction_signals=extract_signals(text),
        )

    def _extract_docling_text(self, conversion_result) -> str:
        document = getattr(conversion_result, "document", None)
        if document is None:
            return ""

        # Docling commonly provides markdown export with table/text layout preserved.
        if hasattr(document, "export_to_markdown"):
            try:
                markdown = document.export_to_markdown()
                if isinstance(markdown, str) and markdown.strip():
                    return markdown
            except Exception:
                pass

        if hasattr(document, "export_to_text"):
            try:
                text = document.export_to_text()
                if isinstance(text, str) and text.strip():
                    return text
            except Exception:
                pass

        return str(document)

    def _parse_image(
        self,
        filename: str,
        raw_bytes: bytes,
        *,
        llm_provider: str = "openai",
        llm_model: str | None = None,
    ) -> ParsedDocument:
        suffix = Path(filename).suffix.lower()
        mime_type = _IMAGE_MIME.get(suffix, "image/jpeg")

        try:
            from PIL import Image as _PILImage
            import io as _io
            img = _PILImage.open(_io.BytesIO(raw_bytes))
            img.verify()
        except ImportError:
            pass  # Pillow not installed — proceed without validation
        except Exception as exc:
            raise ValueError(f"Could not read image file '{filename}': {exc}") from exc

        from backend.app.services.llm_provider import generate_from_image

        try:
            extracted_text = generate_from_image(
                provider=llm_provider,  # type: ignore[arg-type]
                image_bytes=raw_bytes,
                mime_type=mime_type,
                prompt=_IMAGE_EXTRACT_PROMPT,
                model=llm_model,
                temperature=0.0,
            )
        except Exception as exc:
            raise ValueError(f"Image text extraction failed: {exc}") from exc

        if not extracted_text.strip():
            raise ValueError("No text could be extracted from the image.")

        sections = self._sections_from_markdown_or_text(extracted_text, fallback_heading=Path(filename).stem)
        return self._build_document(
            filename=filename,
            file_type="image",
            title=Path(filename).stem,
            text=extracted_text,
            sections=sections,
        )

    def _parse_excel(self, filename: str, raw_bytes: bytes) -> ParsedDocument:
        try:
            import openpyxl
        except ModuleNotFoundError as exc:
            raise RuntimeError(
                "openpyxl is not installed. Run `pip install openpyxl` to enable Excel parsing."
            ) from exc

        from io import BytesIO as _BytesIO
        try:
            workbook = openpyxl.load_workbook(filename=_BytesIO(raw_bytes), data_only=True)
        except Exception as exc:
            raise ValueError(f"Could not open Excel file '{filename}': {exc}") from exc

        sections: list[ParsedSection] = []
        all_text_parts: list[str] = []

        for sheet_name in workbook.sheetnames:
            sheet = workbook[sheet_name]
            rows = list(sheet.iter_rows(values_only=True))
            non_empty_rows = [r for r in rows if any(cell is not None and str(cell).strip() for cell in r)]
            if not non_empty_rows:
                continue

            str_rows: list[list[str]] = [
                [str(cell) if cell is not None else "" for cell in row]
                for row in non_empty_rows
            ]
            max_cols = max(len(r) for r in str_rows)
            padded = [r + [""] * (max_cols - len(r)) for r in str_rows]
            header = padded[0]
            separator = ["---"] * max_cols
            data_rows = padded[1:] if len(padded) > 1 else []

            md_lines: list[str] = []
            md_lines.append("| " + " | ".join(header) + " |")
            md_lines.append("| " + " | ".join(separator) + " |")
            for row in data_rows:
                md_lines.append("| " + " | ".join(row) + " |")

            sheet_text = "\n".join(md_lines)
            sections.append(ParsedSection(heading=sheet_name, body=sheet_text))
            all_text_parts.append("## " + sheet_name + "\n" + sheet_text)

        if not sections:
            raise ValueError(f"Excel file '{filename}' contains no data.")

        full_text = "\n\n".join(all_text_parts)
        return self._build_document(
            filename=filename,
            file_type="xlsx",
            title=Path(filename).stem,
            text=full_text,
            sections=sections,
        )

    def _parse_csv(self, filename: str, raw_bytes: bytes) -> ParsedDocument:
        import csv as _csv

        text = raw_bytes.decode("utf-8", errors="ignore")
        try:
            dialect = _csv.Sniffer().sniff(text[:4096])
        except _csv.Error:
            dialect = _csv.excel  # type: ignore[assignment]

        reader = _csv.reader(text.splitlines(), dialect)
        rows = list(reader)
        rows = [r for r in rows if any(cell.strip() for cell in r)]
        if not rows:
            raise ValueError(f"CSV file '{filename}' contains no data.")

        max_cols = max(len(r) for r in rows)
        padded = [r + [""] * (max_cols - len(r)) for r in rows]
        header = padded[0]
        separator = ["---"] * max_cols
        data_rows = padded[1:] if len(padded) > 1 else []

        md_lines: list[str] = []
        md_lines.append("| " + " | ".join(header) + " |")
        md_lines.append("| " + " | ".join(separator) + " |")
        for row in data_rows:
            md_lines.append("| " + " | ".join(row) + " |")

        md_table = "\n".join(md_lines)
        sections = [ParsedSection(heading=Path(filename).stem, body=md_table)]
        return self._build_document(
            filename=filename,
            file_type="csv",
            title=Path(filename).stem,
            text=md_table,
            sections=sections,
        )

    def _sections_from_markdown_or_text(self, text: str, fallback_heading: str) -> list[ParsedSection]:
        lines = [line.rstrip() for line in text.splitlines()]
        sections: list[ParsedSection] = []
        current_heading = fallback_heading
        current_body: list[str] = []

        for line in lines:
            heading_match = re.match(r"^\s{0,3}#{1,6}\s+(.+)$", line)
            if heading_match:
                if current_body:
                    sections.append(ParsedSection(heading=current_heading, body="\n".join(current_body).strip()))
                    current_body = []
                current_heading = heading_match.group(1).strip()
                continue

            cleaned = line.strip()
            if cleaned:
                current_body.append(cleaned)

        if current_body:
            sections.append(ParsedSection(heading=current_heading, body="\n".join(current_body).strip()))

        if sections:
            return sections

        paragraphs = [chunk.strip() for chunk in text.split("\n\n") if chunk.strip()]
        return [ParsedSection(heading=fallback_heading, body="\n".join(paragraphs))] if paragraphs else []