from __future__ import annotations

from pathlib import Path

from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor, Inches, Emu
from pptx.dml.color import RGBColor as PptxRGBColor
from pptx.util import Inches as PptxInches, Pt as PptxPt, Emu as PptxEmu

from backend.app.config import env


PROJECT_ROOT = Path(__file__).resolve().parents[3]
LOGO_PATH = PROJECT_ROOT / "frontend" / "public" / "bsbi-logo.jpeg"
PPT_TEMPLATE_PATH = PROJECT_ROOT / "backend" / "data" / "bsbi-template.pptx"
PPT_TEMPLATE_LEGACY_PATH = PROJECT_ROOT / "backend" / "data" / "NDA AI Narrative - Show&Tell.pptx"

# ── Company constants (override via env) ───────────────────────────────────
COMPANY_NAME: str = env("COMPANY_NAME", "BSBI Consulting") or "BSBI Consulting"
CONFIDENTIAL_TEXT: str = env("CONFIDENTIAL_TEXT", "CONFIDENTIAL") or "CONFIDENTIAL"

# ── BSBI Brand Colour Palette (docx / shared) ─────────────────────────────
BSBI_RED         = RGBColor(177, 18,  35)
BSBI_DARK_GREY   = RGBColor( 51, 51,  51)
BSBI_MID_GREY    = RGBColor(120, 120, 120)
BSBI_LIGHT_GREY  = RGBColor(230, 230, 230)
BSBI_WHITE       = RGBColor(255, 255, 255)
BSBI_ACCENT_DARK = RGBColor( 40,  40,  40)

# ── PPT-specific colours (pptx RGBColor) ──────────────────────────────────
PPT_RED          = PptxRGBColor(177,  18,  35)   # BSBI primary red
PPT_DARK         = PptxRGBColor( 28,  28,  30)   # near-black for slide backgrounds
PPT_HEADER       = PptxRGBColor( 38,  38,  42)   # dark band on content slides
PPT_ACCENT_DARK  = PptxRGBColor( 28,  28,  30)   # kept for legacy callers
PPT_DARK_GREY    = PptxRGBColor( 51,  51,  51)
PPT_MID_GREY     = PptxRGBColor(120, 120, 120)
PPT_LIGHT_GREY   = PptxRGBColor(230, 230, 230)
PPT_OFF_WHITE    = PptxRGBColor(248, 248, 249)   # slide content background
PPT_WHITE        = PptxRGBColor(255, 255, 255)
PPT_BODY_TEXT    = PptxRGBColor( 36,  36,  36)
PPT_COL_LEFT_BG  = PptxRGBColor(255, 243, 244)   # very light red tint for left col
PPT_COL_RIGHT_BG = PptxRGBColor(245, 245, 247)   # very light grey for right col

# ── Font families ──────────────────────────────────────────────────────────
FONT_HEADING = "Calibri Light"   # all slide/section titles
FONT_BODY    = "Calibri"         # body text, bullets, captions, notes
FONT_FAMILY  = "Calibri"         # docx compatibility alias


def logo_exists() -> bool:
    return LOGO_PATH.exists() and LOGO_PATH.is_file()


def ppt_template_exists() -> bool:
    """Return True if the BSBI template or the legacy NDA template is available."""
    return PPT_TEMPLATE_PATH.exists() or PPT_TEMPLATE_LEGACY_PATH.exists()


def get_ppt_template_path() -> Path | None:
    """Return the best available template path, preferring the generated BSBI one."""
    if PPT_TEMPLATE_PATH.exists():
        return PPT_TEMPLATE_PATH
    if PPT_TEMPLATE_LEGACY_PATH.exists():
        return PPT_TEMPLATE_LEGACY_PATH
    return None


# ── DOCX helpers ───────────────────────────────────────────────────────────

def apply_docx_styles(document) -> None:
    """Configure all shared styles on a python-docx Document."""
    from docx.enum.text import WD_LINE_SPACING

    # Normal
    normal = document.styles["Normal"]
    normal.font.name = FONT_FAMILY
    normal.font.size = Pt(11)
    normal.font.color.rgb = BSBI_DARK_GREY
    pf = normal.paragraph_format
    pf.space_after = Pt(6)
    pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    pf.line_spacing = 1.15

    # Heading 1
    h1 = document.styles["Heading 1"]
    h1.font.name = FONT_FAMILY
    h1.font.size = Pt(16)
    h1.font.bold = True
    h1.font.color.rgb = BSBI_RED
    h1.paragraph_format.space_before = Pt(18)
    h1.paragraph_format.space_after = Pt(6)

    # Heading 2
    h2 = document.styles["Heading 2"]
    h2.font.name = FONT_FAMILY
    h2.font.size = Pt(13)
    h2.font.bold = True
    h2.font.color.rgb = BSBI_DARK_GREY
    h2.paragraph_format.space_before = Pt(12)
    h2.paragraph_format.space_after = Pt(4)

    # List Bullet
    try:
        lb = document.styles["List Bullet"]
        lb.font.name = FONT_FAMILY
        lb.font.size = Pt(11)
        lb.font.color.rgb = BSBI_DARK_GREY
        lb.paragraph_format.space_after = Pt(3)
    except KeyError:
        pass

    for section in document.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)


def add_styled_table(document, headers: list[str], rows: list[list[str]]) -> None:
    """Add a branded table with BSBI red header row and alternating row shading."""
    table = document.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    table.autofit = True

    header_cells = table.rows[0].cells
    for idx, text in enumerate(headers):
        cell = header_cells[idx]
        cell.text = ""
        p = cell.paragraphs[0]
        run = p.add_run(text)
        run.bold = True
        run.font.size = Pt(10)
        run.font.name = FONT_FAMILY
        run.font.color.rgb = BSBI_WHITE
        _shade_cell(cell, "B11223")

    for row_idx, row_data in enumerate(rows):
        row_cells = table.rows[row_idx + 1].cells
        shade = "F5F5F5" if row_idx % 2 == 0 else "FFFFFF"
        for col_idx, text in enumerate(row_data):
            if col_idx >= len(row_cells):
                break
            cell = row_cells[col_idx]
            cell.text = ""
            p = cell.paragraphs[0]
            run = p.add_run(text)
            run.font.size = Pt(10)
            run.font.name = FONT_FAMILY
            run.font.color.rgb = BSBI_DARK_GREY
            _shade_cell(cell, shade)

    document.add_paragraph("")


def _shade_cell(cell, hex_colour: str) -> None:
    from docx.oxml import OxmlElement
    shading_elm = OxmlElement("w:shd")
    shading_elm.set(qn("w:fill"), hex_colour)
    shading_elm.set(qn("w:val"), "clear")
    cell._tc.get_or_add_tcPr().append(shading_elm)


def add_horizontal_rule(document) -> None:
    from docx.oxml import OxmlElement
    p = document.add_paragraph()
    p_pr = p._p.get_or_add_pPr()
    bottom_border = OxmlElement("w:pBdr")
    b = OxmlElement("w:bottom")
    b.set(qn("w:val"), "single")
    b.set(qn("w:sz"), "6")
    b.set(qn("w:space"), "1")
    b.set(qn("w:color"), "B11223")
    bottom_border.append(b)
    p_pr.append(bottom_border)


def add_header_footer(document) -> None:
    from docx.oxml import OxmlElement
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    for section in document.sections:
        section.different_first_page_header_footer = True

        header = section.header
        header.is_linked_to_previous = False
        hp = header.paragraphs[0] if header.paragraphs else header.add_paragraph()
        hp.clear()
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        run = hp.add_run(f"{COMPANY_NAME}  |  {CONFIDENTIAL_TEXT}")
        run.font.size = Pt(8)
        run.font.color.rgb = BSBI_MID_GREY
        run.font.name = FONT_FAMILY

        footer = section.footer
        footer.is_linked_to_previous = False
        fp = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
        fp.clear()
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = fp.add_run()
        run.font.size = Pt(8)
        run.font.color.rgb = BSBI_MID_GREY
        run.font.name = FONT_FAMILY
        fld_char_begin = OxmlElement("w:fldChar")
        fld_char_begin.set(qn("w:fldCharType"), "begin")
        run._r.append(fld_char_begin)
        instr_text = OxmlElement("w:instrText")
        instr_text.set(qn("xml:space"), "preserve")
        instr_text.text = " PAGE "
        run._r.append(instr_text)
        fld_char_end = OxmlElement("w:fldChar")
        fld_char_end.set(qn("w:fldCharType"), "end")
        run._r.append(fld_char_end)


# ── PPT helpers ────────────────────────────────────────────────────────────

def add_ppt_accent_bar(slide, presentation) -> None:
    """Add a thin BSBI-red bar at the very bottom of a slide."""
    slide_width  = presentation.slide_width  or PptxEmu(12192000)
    slide_height = presentation.slide_height or PptxEmu(6858000)
    bar_height = PptxInches(0.22)
    shape = slide.shapes.add_shape(
        1, PptxEmu(0), slide_height - bar_height, slide_width, bar_height,
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = PPT_RED
    shape.line.fill.background()


def add_ppt_slide_number(slide, presentation) -> None:
    """Add a slide-number field in the bottom-right corner."""
    import uuid
    from pptx.enum.text import PP_ALIGN
    from pptx.oxml.ns import qn as pptx_qn
    from lxml import etree

    slide_height = presentation.slide_height or PptxEmu(6858000)
    tb = slide.shapes.add_textbox(
        PptxInches(11.4),
        slide_height - PptxInches(0.38),
        PptxInches(1.1),
        PptxInches(0.28),
    )
    tf = tb.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.RIGHT

    fld = etree.SubElement(p._p, pptx_qn("a:fld"))
    fld.set("id", f"{{{uuid.uuid4()}}}")
    fld.set("type", "slidenum")
    rpr = etree.SubElement(fld, pptx_qn("a:rPr"))
    rpr.set("lang", "en-GB")
    rpr.set("sz", "800")
    solid = etree.SubElement(rpr, pptx_qn("a:solidFill"))
    srgb = etree.SubElement(solid, pptx_qn("a:srgbClr"))
    srgb.set("val", "888888")
    t = etree.SubElement(fld, pptx_qn("a:t"))
    t.text = "‹#›"