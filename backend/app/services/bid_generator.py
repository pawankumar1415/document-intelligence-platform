from __future__ import annotations

import re
from datetime import datetime, timezone

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

from backend.app.models.schemas import GenerateBidRequest, GenerateResult, GeneratedSection
from backend.app.services.branding import (
    BSBI_DARK_GREY,
    BSBI_MID_GREY,
    BSBI_RED,
    FONT_FAMILY,
    LOGO_PATH,
    add_header_footer,
    add_horizontal_rule,
    add_styled_table,
    apply_docx_styles,
    logo_exists,
)
from backend.app.services.file_utils import build_output_path
from backend.app.services.llm_provider import generate_json_object
from backend.app.services.text_utils import build_section_candidates


BID_SYSTEM_PROMPT = """\
You are a senior bid writer at a Big-4 consulting firm drafting a competitive proposal response.

Rules:
- Write in a confident, client-facing, persuasive consulting tone.
- Every paragraph MUST be 2-3 sentences. No filler. No generic platitudes.
- Bullets MUST be specific, concise, and under 20 words each.
- Ground claims in the source material provided. Reference their stated requirements directly.
- The proposal must feel tailored to this specific client opportunity — not a generic template.
- Never write "% improvement" without a real number. Use "To be confirmed during discovery" if no figure is available.
- Return ONLY valid JSON. No markdown, no commentary outside JSON.
"""

BID_USER_PROMPT_TEMPLATE = """\
Return ONLY JSON matching this schema exactly:
{{
  "meta": {{
    "prepared_by": "BSBI Consulting",
    "date": "{today}",
    "version": "Draft 1.0"
  }},
  "sections": [
    {{
      "title": "string",
      "paragraphs": ["string"],
      "bullets": ["string"],
      "table_rows": [["col1", "col2"]]
    }}
  ]
}}

Generate 8 sections in this order:
1. Executive Overview — 2 paragraphs summarising who we are, why we are the right partner, and our understanding of this opportunity. No bullets.
2. Understanding of Your Requirements — 1 paragraph, then 4-6 bullets precisely mirroring the key requirements stated in the source document.
3. Our Proposed Approach — 2 paragraphs on the methodology and how it addresses the stated requirements, 3-4 bullets on key approach principles.
4. Solution Architecture & Capabilities — 2 paragraphs on the technical/functional capabilities we bring, 4-5 bullets on differentiating capabilities.
5. Delivery Plan — 1 paragraph overview, then table_rows with columns ["Phase", "Key Activities", "Duration"] for 3-5 phases. table_rows MUST contain at least 3 rows — never return an empty array.
6. Team & Governance — 1 paragraph on team structure and governance model, 4-5 bullets on key roles or governance principles.
7. Why Choose BSBI — 1 paragraph, then 4-6 bullets on our unique differentiators for this opportunity. Reference the client's stated challenges directly.
8. Proposed Next Steps — 1 short paragraph, then table_rows with columns ["Next Step", "Owner", "Timeline"] for 3-4 action items. table_rows MUST contain at least 3 rows.

Client: {client_name}
Opportunity: {opportunity_title}
Our key strengths to highlight: {our_strengths}

Source content:
{context_blob}
"""


class BidGenerator:
    def generate(self, request: GenerateBidRequest, retrieval_context: list[str] | None = None) -> GenerateResult:
        sections = self._build_sections(request, retrieval_context=retrieval_context or [])
        output_path = build_output_path("bid", "docx")
        self._write_docx(output_path=str(output_path), request=request, sections=sections)
        artifact_name = output_path.name

        summary = f"Bid response drafted for {request.client_name} — {request.opportunity_title}"
        return GenerateResult(
            artifact_type="bid",
            file_path=str(output_path),
            artifact_name=artifact_name,
            download_url=f"/api/v1/artifacts/{artifact_name}",
            summary=summary,
            sections=sections,
        )

    # ── Section building ───────────────────────────────────────────────────

    def _build_sections(self, request: GenerateBidRequest, retrieval_context: list[str]) -> list[GeneratedSection]:
        llm_sections = self._build_sections_from_llm(request, retrieval_context)
        if llm_sections:
            return llm_sections
        return self._build_sections_heuristic(request)

    def _build_sections_from_llm(self, request: GenerateBidRequest, retrieval_context: list[str]) -> list[GeneratedSection] | None:
        today = datetime.now(timezone.utc).strftime("%d %B %Y")
        context_parts = [request.source_document.text[:3000]]
        if retrieval_context:
            context_parts.extend(retrieval_context[:3])
        context_blob = "\n\n---\n\n".join(context_parts)

        prompt = BID_USER_PROMPT_TEMPLATE.format(
            today=today,
            client_name=request.client_name,
            opportunity_title=request.opportunity_title,
            our_strengths="; ".join(request.our_strengths) if request.our_strengths else "To be confirmed",
            context_blob=context_blob,
        )

        try:
            raw = generate_json_object(
                system_prompt=BID_SYSTEM_PROMPT,
                user_prompt=prompt,
                provider=request.llm_provider,
                model=request.llm_model,
            )
        except Exception:
            return None

        sections_raw = raw.get("sections", [])
        if not sections_raw:
            return None

        final_sections: list[GeneratedSection] = []
        for sec in sections_raw:
            if not isinstance(sec, dict):
                continue
            table_rows = list(sec.get("table_rows", []))
            title_lower = sec.get("title", "").lower()

            # Synthesise table_rows from bullets as fallback
            if not table_rows and sec.get("bullets"):
                if any(kw in title_lower for kw in ("delivery", "phase", "plan", "timeline")):
                    table_rows = [
                        [f"Phase {i + 1}", b, "TBC"]
                        for i, b in enumerate(sec["bullets"][:5])
                    ]
                elif any(kw in title_lower for kw in ("next step", "action")):
                    table_rows = [
                        [b, "BSBI / Client", "TBC"]
                        for b in sec["bullets"][:4]
                    ]

            final_sections.append(GeneratedSection(
                title=sec.get("title", ""),
                paragraphs=sec.get("paragraphs", []),
                bullets=sec.get("bullets", []),
                table_rows=table_rows,
            ))

        return final_sections or None

    def _build_sections_heuristic(self, request: GenerateBidRequest) -> list[GeneratedSection]:
        candidates = build_section_candidates(request.source_document)
        return [
            GeneratedSection(
                title="Executive Overview",
                paragraphs=[
                    f"BSBI Consulting is pleased to present this proposal in response to the {request.opportunity_title} opportunity from {request.client_name}.",
                    "We bring deep expertise in delivering complex transformation programmes and a proven track record of results for clients in this sector.",
                ],
            ),
            GeneratedSection(
                title="Understanding of Your Requirements",
                paragraphs=["Based on our review of the source material, we have identified the following key requirements."],
                bullets=self._prepare_lines(candidates.get("Requirements", []) or candidates.get("Objectives", []), limit=6),
            ),
            GeneratedSection(
                title="Our Proposed Approach",
                paragraphs=[
                    "Our approach is structured around three core phases: Discovery, Design & Build, and Deployment.",
                    "Each phase is designed to de-risk delivery while maximising stakeholder engagement and value realisation.",
                ],
                bullets=["Agile delivery with fortnightly sprint reviews", "Dedicated SME engagement from Day 1", "Continuous quality assurance throughout"],
            ),
            GeneratedSection(
                title="Solution Architecture & Capabilities",
                paragraphs=["Our solution leverages proven technologies and frameworks tailored to the stated requirements.", "We have delivered similar architectures at scale across public and private sector clients."],
                bullets=self._prepare_lines(candidates.get("Technical", []), limit=5),
            ),
            GeneratedSection(
                title="Delivery Plan",
                paragraphs=["We propose a phased delivery approach structured to minimise risk and maximise early value."],
                bullets=[],
                table_rows=[
                    ["Phase 1 — Discovery", "Stakeholder workshops, requirements validation, as-is assessment", "4 weeks"],
                    ["Phase 2 — Design & Build", "Solution design, development, integration", "12 weeks"],
                    ["Phase 3 — Deployment", "UAT, training, go-live, hypercare", "4 weeks"],
                ],
            ),
            GeneratedSection(
                title="Team & Governance",
                paragraphs=["Our team will be led by a dedicated Programme Manager with executive sponsorship from a BSBI Partner."],
                bullets=["Programme Manager — single point of accountability", "Weekly steering committee with client stakeholders", "Fortnightly progress reports and risk register reviews"],
            ),
            GeneratedSection(
                title="Why Choose BSBI",
                paragraphs=[f"BSBI Consulting offers {request.client_name} a unique combination of sector knowledge, delivery capability, and commitment to measurable outcomes."],
                bullets=request.our_strengths[:5] if request.our_strengths else ["Deep sector expertise", "Proven delivery at scale", "Fixed-price commercial options available"],
            ),
            GeneratedSection(
                title="Proposed Next Steps",
                paragraphs=["We look forward to discussing this proposal and agreeing a programme of work that delivers real value for your organisation."],
                bullets=[],
                table_rows=[
                    ["Proposal review meeting", "Client & BSBI", "Within 1 week"],
                    ["Clarification Q&A session", "Client & BSBI", "Within 2 weeks"],
                    ["Commercial terms finalised", "Client & BSBI", "Within 3 weeks"],
                ],
            ),
        ]

    @staticmethod
    def _prepare_lines(lines: list[str], limit: int = 5, char_limit: int = 280) -> list[str]:
        result = []
        for line in lines:
            clean = re.sub(r"\s+", " ", line).strip()
            if clean and len(clean) <= char_limit:
                result.append(clean)
            if len(result) >= limit:
                break
        return result

    # ── DOCX writing ───────────────────────────────────────────────────────

    def _write_docx(self, output_path: str, request: GenerateBidRequest, sections: list[GeneratedSection]) -> None:
        doc = Document()
        apply_docx_styles(doc)

        # Cover block
        if logo_exists():
            try:
                doc.add_picture(str(LOGO_PATH), width=Pt(90))
                doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.LEFT
            except Exception:
                pass

        title_para = doc.add_paragraph()
        title_para.alignment = WD_ALIGN_PARAGRAPH.LEFT
        run = title_para.add_run("Proposal Response")
        run.bold = True
        run.font.size = Pt(26)
        run.font.color.rgb = BSBI_RED
        run.font.name = FONT_FAMILY

        opp_para = doc.add_paragraph()
        opp_run = opp_para.add_run(request.opportunity_title)
        opp_run.bold = True
        opp_run.font.size = Pt(16)
        opp_run.font.color.rgb = BSBI_DARK_GREY
        opp_run.font.name = FONT_FAMILY

        client_para = doc.add_paragraph()
        client_run = client_para.add_run(f"Prepared for: {request.client_name}")
        client_run.font.size = Pt(12)
        client_run.font.color.rgb = BSBI_DARK_GREY
        client_run.font.name = FONT_FAMILY

        date_para = doc.add_paragraph()
        date_run = date_para.add_run(datetime.now(timezone.utc).strftime("%d %B %Y"))
        date_run.font.size = Pt(10)
        date_run.font.color.rgb = BSBI_MID_GREY
        date_run.font.name = FONT_FAMILY

        add_horizontal_rule(doc)
        doc.add_paragraph("")

        # Sections
        for section in sections:
            doc.add_heading(section.title, level=1)
            add_horizontal_rule(doc)

            for para in section.paragraphs:
                if para.strip():
                    p = doc.add_paragraph(para.strip())
                    p.style = doc.styles["Normal"]

            for bullet in section.bullets:
                if bullet.strip():
                    doc.add_paragraph(bullet.strip(), style="List Bullet")

            if section.table_rows:
                col_count = max(len(row) for row in section.table_rows) if section.table_rows else 2
                # Build generic column headers based on common patterns
                headers = _infer_table_headers(section.title, col_count)
                add_styled_table(doc, headers=headers, rows=section.table_rows)

            doc.add_paragraph("")

        add_header_footer(doc)
        doc.save(output_path)


def _infer_table_headers(section_title: str, col_count: int) -> list[str]:
    title_lower = section_title.lower()
    if col_count == 3:
        if any(kw in title_lower for kw in ("delivery", "phase", "plan")):
            return ["Phase", "Key Activities", "Duration"]
        if any(kw in title_lower for kw in ("next step", "action")):
            return ["Next Step", "Owner", "Timeline"]
    if col_count == 2:
        return ["Item", "Details"]
    return [f"Column {i + 1}" for i in range(col_count)]