"""
rules_parser.py — Extract compliance criteria from uploaded rules documents.

Supports .docx, .pdf, .xlsx, .xls, .csv, .txt
Strategy:
  - Excel/CSV  → direct column detection (name, description, severity)
  - DOCX/PDF   → extract full text → LLM extraction → fallback to regex patterns
"""
from __future__ import annotations

import io
import logging
import re
from typing import Any

logger = logging.getLogger(__name__)

_HIGH_KEYWORDS = {"must", "shall", "required", "mandatory", "critical", "essential"}
_LOW_KEYWORDS  = {"may", "optional", "consider", "suggested", "recommended"}


def _infer_severity(text: str) -> str:
    lower = text.lower()
    for kw in _HIGH_KEYWORDS:
        if re.search(rf"\b{kw}\b", lower):
            return "high"
    for kw in _LOW_KEYWORDS:
        if re.search(rf"\b{kw}\b", lower):
            return "low"
    return "medium"


# ── Excel / CSV ────────────────────────────────────────────────────────────────

def _parse_excel_rules(raw_bytes: bytes, filename: str) -> list[dict]:
    import pandas as pd

    ext = filename.lower().rsplit(".", 1)[-1]
    df = pd.read_csv(io.BytesIO(raw_bytes), dtype=str) if ext == "csv" \
        else pd.read_excel(io.BytesIO(raw_bytes), dtype=str)
    df.fillna("", inplace=True)

    col_lower = {c.lower().strip(): c for c in df.columns}

    name_col = next(
        (col_lower[k] for k in ("criterion", "rule", "name", "title", "requirement", "criteria") if k in col_lower),
        df.columns[0] if len(df.columns) > 0 else None,
    )
    desc_col = next(
        (col_lower[k] for k in ("description", "detail", "explanation", "text", "guidance") if k in col_lower),
        None,
    )
    sev_col = next(
        (col_lower[k] for k in ("severity", "priority", "level", "weight") if k in col_lower),
        None,
    )

    criteria = []
    for _, row in df.iterrows():
        name = str(row.get(name_col, "")).strip() if name_col else ""
        if not name or name.lower() in ("nan", "none", ""):
            continue
        description = str(row.get(desc_col, "")).strip() if desc_col else ""
        sev_raw = str(row.get(sev_col, "")).strip().lower() if sev_col else ""
        severity = sev_raw if sev_raw in ("high", "medium", "low") else _infer_severity(f"{name} {description}")
        criteria.append({"name": name[:120], "description": description or name, "severity": severity})
    return criteria


# ── Text extraction ────────────────────────────────────────────────────────────

def _extract_docx_text(raw_bytes: bytes) -> str:
    from docx import Document
    doc = Document(io.BytesIO(raw_bytes))
    parts: list[str] = []
    for para in doc.paragraphs:
        t = para.text.strip()
        if t:
            parts.append(t)
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def _extract_pdf_text(raw_bytes: bytes) -> str:
    import pdfplumber
    parts: list[str] = []
    with pdfplumber.open(io.BytesIO(raw_bytes)) as pdf:
        for page in pdf.pages:
            t = page.extract_text() or ""
            if t.strip():
                parts.append(t.strip())
    return "\n".join(parts)


# ── LLM extraction ─────────────────────────────────────────────────────────────

def _extract_with_llm(text: str, provider: str, model: str | None) -> list[dict]:
    from backend.app.services.llm_provider import generate_json_object

    system_prompt = """\
You are a compliance document analyst. Extract every compliance criterion or rule from the text.
For each, return: name (≤10 words), description (full wording), severity (high/medium/low).

Severity rules:
  high   — uses "must", "shall", "required", "mandatory", "critical"
  low    — uses "may", "optional", "consider", "suggested"
  medium — everything else

Return ONLY valid JSON:
{"criteria": [{"name": "...", "description": "...", "severity": "high|medium|low"}]}"""

    try:
        result = generate_json_object(
            provider=provider,
            system_prompt=system_prompt,
            user_prompt=f"DOCUMENT:\n{text[:6000]}",
            temperature=0.1,
            model=model,
        )
        items = result.get("criteria", [])
        criteria = []
        for item in items:
            if not isinstance(item, dict):
                continue
            name = str(item.get("name", "")).strip()
            if not name:
                continue
            sev = str(item.get("severity", "medium")).lower()
            if sev not in ("high", "medium", "low"):
                sev = "medium"
            criteria.append({
                "name": name[:120],
                "description": str(item.get("description", name)).strip(),
                "severity": sev,
            })
        return criteria
    except Exception as exc:
        logger.warning("LLM rules extraction failed: %s", exc)
        return []


# ── Regex fallback ─────────────────────────────────────────────────────────────

def _regex_parse(text: str) -> list[dict]:
    criteria: list[dict] = []
    for pattern in (r"^\s*\d+[\.\)]\s+(.+)$", r"^\s*[-•*]\s+(.+)$"):
        matches = re.findall(pattern, text, re.MULTILINE)
        for item in matches:
            item = item.strip()
            if len(item) > 5:
                criteria.append({
                    "name": item[:120],
                    "description": item,
                    "severity": _infer_severity(item),
                })
        if criteria:
            return criteria
    return criteria


# ── Public API ─────────────────────────────────────────────────────────────────

def parse_rules_document(
    raw_bytes: bytes,
    filename: str,
    provider: str = "ollama",
    model: str | None = None,
) -> list[dict[str, Any]]:
    """
    Parse a rules/compliance document and return criteria dicts.
    Each dict: {name, description, severity}
    """
    ext = filename.lower().rsplit(".", 1)[-1]

    if ext in ("xlsx", "xls", "csv"):
        return _parse_excel_rules(raw_bytes, filename)

    text = ""
    if ext == "docx":
        try:
            text = _extract_docx_text(raw_bytes)
        except Exception as exc:
            logger.warning("DOCX extraction failed: %s", exc)
    elif ext == "pdf":
        try:
            text = _extract_pdf_text(raw_bytes)
        except Exception as exc:
            logger.warning("PDF extraction failed: %s", exc)
    else:
        text = raw_bytes.decode("utf-8", errors="replace")

    if not text.strip():
        return []

    criteria = _extract_with_llm(text, provider, model)
    if not criteria:
        criteria = _regex_parse(text)
    if not criteria:
        criteria = [{"name": "Compliance Requirement", "description": text[:500], "severity": "high"}]

    return criteria