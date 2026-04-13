from __future__ import annotations

import re
from datetime import datetime, timezone

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.shared import Inches, Pt, RGBColor

from backend.app.models.schemas import GenerateCaseStudyRequest, GenerateResult, GeneratedSection
from backend.app.services.branding import (
    BSBI_DARK_GREY,
    BSBI_MID_GREY,
    BSBI_RED,
    FONT_FAMILY,
    LOGO_PATH,
    add_header_footer,
    add_horizontal_rule,
    apply_docx_styles,
    logo_exists,
)
from backend.app.services.file_utils import build_output_path
from backend.app.services.llm_provider import generate_json_object


_SYSTEM_PROMPT = """\
You are a senior consultant writing a polished client-facing case study for a consulting firm.

Rules:
- Write in third person, past tense for approach/delivery, present tense for outcomes.
- Be specific. Reference actual details from the source material — names, figures, timelines.
- Keep paragraphs to 2-3 sentences. Bullets must be under 20 words each.
- Do not use generic filler. Every sentence must add value.
- If a specific metric is not in the source material, write a believable qualitative outcome instead.
- Return ONLY valid JSON. No markdown or commentary outside the JSON.
"""

_USER_PROMPT_TEMPLATE = """\
Write a consulting case study using the source content below.

Return ONLY JSON matching this schema:
{{
  "background": ["paragraph1", "paragraph2"],
  "challenge": {{
    "summary": "1-2 sentence challenge statement",
    "bullets": ["specific challenge bullet 1", "specific challenge bullet 2", "specific challenge bullet 3", "specific challenge bullet 4"]
  }},
  "approach": {{
    "summary": "1-2 sentence approach overview",
    "steps": ["step 1 description", "step 2 description", "step 3 description", "step 4 description", "step 5 description"]
  }},
  "results": {{
    "summary": "1-2 sentence results narrative",
    "bullets": ["outcome bullet 1", "outcome bullet 2", "outcome bullet 3", "outcome bullet 4"]
  }},
  "outcome": ["paragraph1", "paragraph2"]
}}

Client: {client_name}
Industry: {client_industry}
Engagement: {engagement_title}
Challenge context: {challenge_summary}
Key approach points: {approach_points}

Source content:
{source_text}
"""


class CaseStudyGenerator:
    def generate(self, request: GenerateCaseStudyRequest) -> GenerateResult:
        sections = self._build_sections(request)
        output_path = build_output_path("case-study", "docx")
        self._write_docx(output_path=str(output_path), request=request, sections=sections)
        artifact_name = output_path.name

        summary = f"Case study generated for {request.client_name} — {request.engagement_title}"
        return GenerateResult(
            artifact_type="case_study",
            file_path=str(output_path),
            artifact_name=artifact_name,
            download_url=f"/api/v1/artifacts/{artifact_name}",
            summary=summary,
            sections=sections,
        )

    # ── Section building ───────────────────────────────────────────────────

    def _build_sections(self, request: GenerateCaseStudyRequest) -> list[GeneratedSection]:
        llm = self._build_from_llm(request)
        if llm:
            return llm
        return self._build_fallback(request)

    def _build_from_llm(self, request: GenerateCaseStudyRequest) -> list[GeneratedSection] | None:
        prompt = _USER_PROMPT_TEMPLATE.format(
            client_name=request.client_name,
            client_industry=request.client_industry or "Consulting",
            engagement_title=request.engagement_title,
            challenge_summary=request.challenge_summary or "To be derived from source content",
            approach_points="; ".join(request.our_approach_points) if request.our_approach_points else "To be derived from source content",
            source_text=request.source_document.text[:5000],
        )

        try:
            raw = generate_json_object(
                system_prompt=_SYSTEM_PROMPT,
                user_prompt=prompt,
                provider=request.llm_provider,
                model=request.llm_model,
            )
        except Exception:
            return None

        if not isinstance(raw, dict):
            return None

        sections: list[GeneratedSection] = []

        bg = raw.get("background", [])
        if bg:
            sections.append(GeneratedSection(
                title="Background",
                paragraphs=[p for p in bg if isinstance(p, str) and p.strip()],
            ))

        ch = raw.get("challenge", {})
        if ch:
            bullets = [b for b in ch.get("bullets", []) if isinstance(b, str)]
            sections.append(GeneratedSection(
                title="The Challenge",
                paragraphs=[ch.get("summary", "")] if ch.get("summary") else [],
                bullets=bullets,
            ))

        ap = raw.get("approach", {})
        if ap:
            steps = [s for s in ap.get("steps", []) if isinstance(s, str)]
            sections.append(GeneratedSection(
                title="Our Approach",
                paragraphs=[ap.get("summary", "")] if ap.get("summary") else [],
                bullets=steps,
            ))

        re_data = raw.get("results", {})
        if re_data:
            bullets = [b for b in re_data.get("bullets", []) if isinstance(b, str)]
            sections.append(GeneratedSection(
                title="Results",
                paragraphs=[re_data.get("summary", "")] if re_data.get("summary") else [],
                bullets=bullets,
            ))

        oc = raw.get("outcome", [])
        if oc:
            sections.append(GeneratedSection(
                title="Outcome & Impact",
                paragraphs=[p for p in oc if isinstance(p, str) and p.strip()],
            ))

        return sections if sections else None

    def _build_fallback(self, request: GenerateCaseStudyRequest) -> list[GeneratedSection]:
        return [
            GeneratedSection(
                title="Background",
                paragraphs=[
                    f"{request.client_name} engaged BSBI Consulting to support their {request.engagement_title} programme.",
                    "The engagement was designed to deliver measurable outcomes aligned to the organisation's strategic priorities.",
                ],
            ),
            GeneratedSection(
                title="The Challenge",
                paragraphs=[request.challenge_summary or "The client faced a complex operational challenge requiring specialist expertise."],
                bullets=["Fragmented processes and lack of integrated data", "Limited internal capacity for programme delivery", "Tight regulatory and compliance requirements"],
            ),
            GeneratedSection(
                title="Our Approach",
                paragraphs=["BSBI deployed a structured methodology tailored to the client's context and priorities."],
                bullets=request.our_approach_points[:5] if request.our_approach_points else [
                    "Discovery and stakeholder engagement", "Current-state assessment and gap analysis",
                    "Solution design and recommendations", "Delivery support and quality assurance",
                ],
            ),
            GeneratedSection(
                title="Results",
                paragraphs=["The engagement delivered tangible outcomes across all workstreams."],
                bullets=["Improved operational efficiency and reduced manual effort", "Clear governance and accountability framework established", "Stakeholder confidence increased"],
            ),
            GeneratedSection(
                title="Outcome & Impact",
                paragraphs=[
                    f"{request.client_name} now has a robust foundation for sustainable improvement.",
                    "BSBI Consulting continues to support the organisation as a trusted delivery partner.",
                ],
            ),
        ]

    # ── DOCX writing ───────────────────────────────────────────────────────

    def _write_docx(self, output_path: str, request: GenerateCaseStudyRequest, sections: list[GeneratedSection]) -> None:
        doc = Document()
        apply_docx_styles(doc)

        # ── Cover banner ──────────────────────────────────────────────────
        self._add_cover(doc, request)
        add_horizontal_rule(doc)
        doc.add_paragraph("")

        # ── Metrics strip (if any) ────────────────────────────────────────
        if request.headline_metrics:
            self._add_metrics_strip(doc, request)
            doc.add_paragraph("")

        # ── Sections ─────────────────────────────────────────────────────
        for i, section in enumerate(sections):
            self._add_section(doc, section, is_challenge=(i == 1))
            doc.add_paragraph("")

        add_header_footer(doc)
        doc.save(output_path)

    def _add_cover(self, doc: Document, request: GenerateCaseStudyRequest) -> None:
        if logo_exists():
            try:
                doc.add_picture(str(LOGO_PATH), width=Pt(90))
                doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.LEFT
            except Exception:
                pass

        # "CASE STUDY" label
        tag_para = doc.add_paragraph()
        tag_run = tag_para.add_run("CASE STUDY")
        tag_run.bold = True
        tag_run.font.size = Pt(11)
        tag_run.font.color.rgb = BSBI_RED
        tag_run.font.name = FONT_FAMILY

        # Engagement title
        title_para = doc.add_paragraph()
        title_run = title_para.add_run(request.engagement_title)
        title_run.bold = True
        title_run.font.size = Pt(22)
        title_run.font.color.rgb = BSBI_DARK_GREY
        title_run.font.name = FONT_FAMILY

        # Client name
        client_para = doc.add_paragraph()
        client_run = client_para.add_run(request.client_name)
        client_run.font.size = Pt(14)
        client_run.font.color.rgb = BSBI_RED
        client_run.font.name = FONT_FAMILY
        client_run.bold = True

        # Industry + date
        meta_parts = []
        if request.client_industry:
            meta_parts.append(request.client_industry)
        meta_parts.append(datetime.now(timezone.utc).strftime("%B %Y"))
        meta_para = doc.add_paragraph()
        meta_run = meta_para.add_run("  ·  ".join(meta_parts))
        meta_run.font.size = Pt(10)
        meta_run.font.color.rgb = BSBI_MID_GREY
        meta_run.font.name = FONT_FAMILY

    def _add_metrics_strip(self, doc: Document, request: GenerateCaseStudyRequest) -> None:
        metrics = request.headline_metrics[:4]  # max 4 per row
        col_count = len(metrics)
        table = doc.add_table(rows=1, cols=col_count)
        table.style = "Table Grid"
        for i, metric in enumerate(metrics):
            cell = table.cell(0, i)
            # Red background
            tc = cell._tc
            tcPr = tc.get_or_add_tcPr()
            shd = OxmlElement("w:shd")
            shd.set(qn("w:val"), "clear")
            shd.set(qn("w:color"), "auto")
            shd.set(qn("w:fill"), "B11223")
            tcPr.append(shd)

            # Value (big)
            val_para = cell.paragraphs[0]
            val_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            val_run = val_para.add_run(metric.value)
            val_run.bold = True
            val_run.font.size = Pt(24)
            val_run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
            val_run.font.name = FONT_FAMILY

            # Label
            lbl_para = cell.add_paragraph()
            lbl_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            lbl_run = lbl_para.add_run(metric.label)
            lbl_run.bold = True
            lbl_run.font.size = Pt(9)
            lbl_run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
            lbl_run.font.name = FONT_FAMILY

            # Description (if any)
            if metric.description:
                desc_para = cell.add_paragraph()
                desc_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
                desc_run = desc_para.add_run(metric.description)
                desc_run.font.size = Pt(8)
                desc_run.font.color.rgb = RGBColor(0xFF, 0xCC, 0xCC)
                desc_run.font.name = FONT_FAMILY

    def _add_section(self, doc: Document, section: GeneratedSection, is_challenge: bool = False) -> None:
        # Section heading
        heading = doc.add_heading(section.title, level=1)
        add_horizontal_rule(doc)

        # For the Challenge section, add an accent-styled block
        if is_challenge and section.paragraphs:
            # Use a table cell for the colored box effect
            tbl = doc.add_table(rows=1, cols=1)
            tbl.style = "Table Grid"
            cell = tbl.cell(0, 0)
            tc = cell._tc
            tcPr = tc.get_or_add_tcPr()
            shd = OxmlElement("w:shd")
            shd.set(qn("w:val"), "clear")
            shd.set(qn("w:color"), "auto")
            shd.set(qn("w:fill"), "F5E6E8")
            tcPr.append(shd)
            for para_text in section.paragraphs:
                if para_text.strip():
                    p = cell.add_paragraph(para_text.strip())
                    for run in p.runs:
                        run.font.size = Pt(11)
                        run.font.name = FONT_FAMILY
            # Remove the first empty paragraph added by default
            if cell.paragraphs and not cell.paragraphs[0].text:
                p_elem = cell.paragraphs[0]._element
                p_elem.getparent().remove(p_elem)
            doc.add_paragraph("")
        else:
            for para_text in section.paragraphs:
                if para_text.strip():
                    p = doc.add_paragraph(para_text.strip())
                    p.style = doc.styles["Normal"]

        # Numbered approach steps vs plain bullets
        if section.title == "Our Approach":
            for idx, bullet in enumerate(section.bullets, 1):
                if bullet.strip():
                    p = doc.add_paragraph()
                    p.style = doc.styles["Normal"]
                    num_run = p.add_run(f"{idx:02d}  ")
                    num_run.bold = True
                    num_run.font.color.rgb = BSBI_RED
                    num_run.font.name = FONT_FAMILY
                    text_run = p.add_run(bullet.strip())
                    text_run.font.name = FONT_FAMILY
        else:
            for bullet in section.bullets:
                if bullet.strip():
                    doc.add_paragraph(bullet.strip(), style="List Bullet")