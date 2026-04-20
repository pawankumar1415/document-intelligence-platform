from __future__ import annotations

import logging
from datetime import datetime, timezone
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

from backend.app.models.schemas import ExtractionResult, ExtractRequest
from backend.app.services import persistence
from backend.app.services.branding import (
    BSBI_DARK_GREY,
    BSBI_RED,
    FONT_FAMILY,
    LOGO_PATH,
    add_header_footer,
    add_styled_table,
    apply_docx_styles,
    logo_exists,
)
from backend.app.services.file_utils import OUTPUT_DIR, build_output_path
from backend.app.services.llm_provider import generate_json_object

logger = logging.getLogger(__name__)

_SYSTEM_PROMPT = """\
You are a precise document analyst. Extract structured entities from the provided document text.
You will be given an entity type and a list of fields to extract for each entity.

Rules:
- Extract ONLY entities that are explicitly present in the document. Do not invent or infer.
- If a field's value is not stated, use an empty string "".
- Return as many entities as you find — be thorough.
- Return ONLY valid JSON — no markdown, no commentary outside the JSON.
"""

_USER_PROMPT = """\
Extract all "{entity_label}" entities from the document below.

For each entity, extract these fields:
{fields_description}

Return JSON matching this schema exactly:
{{
  "entities": [
    {{
      {field_schema}
    }}
  ]
}}

If no entities are found, return: {{"entities": []}}

Document — "{document_name}":
{text}
"""


def _build_field_schema(fields: list[dict]) -> str:
    return ",\n      ".join(f'"{f["name"]}": "string"' for f in fields)


def _build_fields_description(fields: list[dict]) -> str:
    lines = []
    for f in fields:
        req = "required" if f.get("required", True) else "optional"
        desc = f.get("description", "")
        lines.append(f'- {f["name"]} ({req}): {desc}')
    return "\n".join(lines)


def run_extraction(request: ExtractRequest, user_id: int) -> ExtractionResult:
    schema = persistence.get_extraction_schema(request.schema_id, user_id)
    if not schema:
        raise ValueError(f"Extraction schema {request.schema_id} not found.")

    fields: list[dict] = schema.get("fields", [])
    entity_label: str = schema["entity_label"]
    field_names = [f["name"] for f in fields]

    prompt = _USER_PROMPT.format(
        entity_label=entity_label,
        fields_description=_build_fields_description(fields),
        field_schema=_build_field_schema(fields),
        document_name=request.document_name,
        text=request.text[:5000],
    )

    try:
        raw = generate_json_object(
            system_prompt=_SYSTEM_PROMPT,
            user_prompt=prompt,
            provider=request.llm_provider,
            model=request.llm_model,
        )
    except Exception as exc:
        logger.warning("LLM extraction failed: %s", exc)
        raw = {"entities": []}

    entities: list[dict] = raw.get("entities", []) if isinstance(raw.get("entities"), list) else []

    # Build row data aligned with field_names
    rows: list[list[str]] = []
    for entity in entities:
        if not isinstance(entity, dict):
            continue
        row = [str(entity.get(fn, "")) for fn in field_names]
        rows.append(row)

    output_path = _write_register_docx(
        schema_name=schema["name"],
        entity_label=entity_label,
        document_name=request.document_name,
        field_names=field_names,
        rows=rows,
    )

    artifact_name = output_path.name
    return ExtractionResult(
        schema_name=schema["name"],
        entity_label=entity_label,
        field_names=field_names,
        rows=rows,
        artifact_name=artifact_name,
        download_url=f"/api/v1/artifacts/{artifact_name}",
        total_extracted=len(rows),
    )


def _write_register_docx(
    schema_name: str,
    entity_label: str,
    document_name: str,
    field_names: list[str],
    rows: list[list[str]],
) -> Path:
    doc = Document()
    apply_docx_styles(doc)

    # Cover / title block
    if logo_exists():
        try:
            doc.add_picture(str(LOGO_PATH), width=Pt(90))
            doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.LEFT
        except Exception:
            pass

    title_para = doc.add_paragraph()
    title_para.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = title_para.add_run(f"{schema_name} Register")
    run.bold = True
    run.font.size = Pt(22)
    run.font.color.rgb = BSBI_RED
    run.font.name = FONT_FAMILY

    subtitle = doc.add_paragraph()
    run2 = subtitle.add_run(f"Source Document: {document_name}")
    run2.font.size = Pt(11)
    run2.font.color.rgb = BSBI_DARK_GREY
    run2.font.name = FONT_FAMILY

    date_str = datetime.now(timezone.utc).strftime("%d %B %Y")
    date_para = doc.add_paragraph()
    run3 = date_para.add_run(f"Generated: {date_str}")
    run3.font.size = Pt(10)
    run3.font.color.rgb = BSBI_DARK_GREY
    run3.font.name = FONT_FAMILY

    doc.add_paragraph("")

    summary_para = doc.add_paragraph()
    summary_run = summary_para.add_run(f"Total {entity_label} entries extracted: {len(rows)}")
    summary_run.bold = True
    summary_run.font.size = Pt(11)
    summary_run.font.color.rgb = BSBI_DARK_GREY
    summary_run.font.name = FONT_FAMILY

    doc.add_paragraph("")

    if rows:
        add_styled_table(doc, headers=field_names, rows=rows)
    else:
        empty_para = doc.add_paragraph()
        empty_run = empty_para.add_run(f"No {entity_label} entities were identified in the source document.")
        empty_run.italic = True
        empty_run.font.color.rgb = BSBI_DARK_GREY
        empty_run.font.name = FONT_FAMILY

    add_header_footer(doc)

    timestamp = datetime.now(timezone.utc).strftime("%Y%m%d%H%M%S")
    output_path = OUTPUT_DIR / f"register-{timestamp}.docx"
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    doc.save(str(output_path))
    return output_path